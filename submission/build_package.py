#!/usr/bin/env python3
"""Build the Knowledge-Based Systems upload package from the manuscript."""

from __future__ import annotations

import csv
import json
import re
import shutil
import subprocess
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SUB = Path(__file__).resolve().parent
MANUSCRIPT = ROOT / "MANUSCRIPT.md"

HIGHLIGHTS = [
    "Scores closure from the first priced round and blocks lifetime funding fields.",
    "Four of seven certainty-factor rules survive a pre-test retention check.",
    "Top 30% of booster scores: precision 0.760 against a base rate of 0.510.",
    "Logistic regression matches the booster; stacking does not lift the queue.",
    "Fired rules stay on the warning card and do not enter the probability.",
]


def ensure_dirs() -> dict[str, Path]:
    names = [
        "01_manuscript",
        "02_figures",
        "03_tables",
        "04_supporting_data",
        "05_graphical_abstract",
        "06_supplementary",
        "07_cover_letter_and_highlights",
    ]
    paths = {}
    for name in names:
        path = SUB / name
        path.mkdir(parents=True, exist_ok=True)
        paths[name] = path
    return paths


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text.replace("`", ""))
    set_run_font(run, **kwargs)
    return run


def double_space(paragraph):
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(0)


def add_line_numbers(section):
    sect_pr = section._sectPr
    numbering = OxmlElement("w:lnNumType")
    numbering.set(qn("w:countBy"), "1")
    numbering.set(qn("w:restart"), "continuous")
    sect_pr.append(numbering)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def parse_markdown(text: str) -> list[tuple]:
    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|\s*-+", lines[i]):
                    rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(("table", rows))
            continue
        image = re.match(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if image:
            blocks.append(("image", image.group(2), image.group(1)))
            i += 1
            continue
        if line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
            i += 1
            continue
        if line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
            i += 1
            continue
        if line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        parts = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "!")):
            parts.append(lines[i].strip())
            i += 1
        blocks.append(("p", " ".join(parts)))
    return blocks


def new_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    add_line_numbers(section)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "KB-EWS short communication  |  Knowledge-Based Systems", size=9)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return doc


def add_paragraph(doc, text, *, size=12, bold=False, italic=False, center=False, space_after=0):
    paragraph = doc.add_paragraph()
    double_space(paragraph)
    paragraph.paragraph_format.space_after = Pt(space_after)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(paragraph, text, size=size, bold=bold, italic=italic)
    return paragraph


def add_table(doc, rows: list[list[str]]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_text(paragraph, value, size=9, bold=(r == 0))
            if r == 0:
                shade_cell(cell, "E7EEF5")
    doc.add_paragraph()


def build_manuscript_docx(blocks, path: Path):
    doc = new_doc()
    title = next(text for kind, text in blocks if kind == "h1")
    add_paragraph(doc, title, size=16, bold=True, center=True, space_after=6)
    add_paragraph(doc, "Short Communication", size=12, italic=True, center=True)
    add_paragraph(doc, "Balavaraprasad Reddy Yenugula", size=12, bold=True, center=True)
    add_paragraph(doc, "Doctor of Business Administration candidate", size=12, italic=True, center=True, space_after=12)
    add_paragraph(
        doc,
        "Submitted to Knowledge-Based Systems (Elsevier). Article type: Short Communication.",
        size=11,
        italic=True,
        center=True,
        space_after=12,
    )
    started = False
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            continue
        if kind == "p" and block[1] in {
            "Short Communication",
            "Balavaraprasad Reddy Yenugula",
            "Doctor of Business Administration candidate",
        }:
            continue
        if kind == "h2" and not started:
            started = True
        if kind == "h2":
            add_paragraph(doc, block[1], size=14, bold=True, space_after=6)
        elif kind == "h3":
            add_paragraph(doc, block[1], size=12, bold=True, italic=True)
        elif kind == "p":
            if block[1].startswith("Keywords:"):
                paragraph = doc.add_paragraph()
                double_space(paragraph)
                add_text(paragraph, "Keywords: ", size=12, bold=True)
                add_text(paragraph, block[1].split(":", 1)[1].strip(), size=12, italic=True)
                add_paragraph(doc, "Highlights", size=14, bold=True, space_after=6)
                for item in HIGHLIGHTS:
                    paragraph = doc.add_paragraph()
                    double_space(paragraph)
                    add_text(paragraph, "• " + item, size=12)
            else:
                add_paragraph(doc, block[1])
        elif kind == "table":
            add_table(doc, block[1])
        elif kind == "image":
            image_path = ROOT / block[1]
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(6.2))
    doc.save(path)


def html_escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("`", "")
    )


def build_manuscript_html(blocks, path: Path):
    parts = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'><title>KB-EWS</title>",
        "<style>",
        "@page { size: letter; margin: 1in; }",
        "body { font-family: 'Times New Roman', Times, serif; font-size: 12pt; line-height: 2; color: #111; }",
        "h1 { font-size: 16pt; text-align: center; line-height: 1.3; margin-bottom: 0.2em; }",
        "h2 { font-size: 14pt; line-height: 1.4; margin-top: 1.2em; }",
        "h3 { font-size: 12pt; font-style: italic; line-height: 1.4; }",
        ".center { text-align: center; }",
        ".note { font-style: italic; text-align: center; }",
        "table { border-collapse: collapse; width: 100%; font-size: 9pt; line-height: 1.3; margin: 0.6em 0 1em; }",
        "th, td { border: 1px solid #333; padding: 4px 6px; vertical-align: top; }",
        "th { background: #E7EEF5; }",
        "img { width: 100%; max-width: 640px; display: block; margin: 0.4em auto; }",
        "figcaption, .caption { font-size: 10pt; line-height: 1.35; margin: 0 0 1em; }",
        "ul { margin-top: 0; }",
        "</style></head><body>",
        f"<h1>{html_escape(next(text for kind, text in blocks if kind == 'h1'))}</h1>",
        "<p class='center'><em>Short Communication</em></p>",
        "<p class='center'><strong>Balavaraprasad Reddy Yenugula</strong></p>",
        "<p class='center'><em>Doctor of Business Administration candidate</em></p>",
        "<p class='note'>Submitted to Knowledge-Based Systems (Elsevier). Article type: Short Communication.</p>",
    ]
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            continue
        if kind == "p" and block[1] in {
            "Short Communication",
            "Balavaraprasad Reddy Yenugula",
            "Doctor of Business Administration candidate",
        }:
            continue
        if kind == "h2":
            parts.append(f"<h2>{html_escape(block[1])}</h2>")
        elif kind == "h3":
            parts.append(f"<h3>{html_escape(block[1])}</h3>")
        elif kind == "p":
            if block[1].startswith("Keywords:"):
                parts.append(f"<p><strong>Keywords:</strong> <em>{html_escape(block[1].split(':', 1)[1].strip())}</em></p>")
                parts.append("<h2>Highlights</h2><ul>")
                for item in HIGHLIGHTS:
                    parts.append(f"<li>{html_escape(item)}</li>")
                parts.append("</ul>")
            elif block[1].startswith("Figure "):
                parts.append(f"<p class='caption'><strong>{html_escape(block[1])}</strong></p>")
            else:
                parts.append(f"<p>{html_escape(block[1])}</p>")
        elif kind == "table":
            parts.append("<table>")
            for r, row in enumerate(block[1]):
                tag = "th" if r == 0 else "td"
                parts.append("<tr>" + "".join(f"<{tag}>{html_escape(cell)}</{tag}>" for cell in row) + "</tr>")
            parts.append("</table>")
        elif kind == "image":
            src = (ROOT / block[1]).resolve().as_uri()
            parts.append(f"<img src='{src}' alt='{html_escape(block[2])}'>")
    parts.append("</body></html>")
    path.write_text("\n".join(parts), encoding="utf-8")


def write_simple_docx(path: Path, title: str, paragraphs: list[str]):
    doc = new_doc()
    add_paragraph(doc, title, size=16, bold=True, center=True, space_after=12)
    for paragraph in paragraphs:
        if paragraph.startswith("# "):
            add_paragraph(doc, paragraph[2:], size=14, bold=True, space_after=6)
        else:
            add_paragraph(doc, paragraph)
    doc.save(path)


def write_simple_html(path: Path, title: str, paragraphs: list[str]):
    body = [f"<h1>{html_escape(title)}</h1>"]
    for paragraph in paragraphs:
        if paragraph.startswith("# "):
            body.append(f"<h2>{html_escape(paragraph[2:])}</h2>")
        else:
            body.append(f"<p>{html_escape(paragraph)}</p>")
    path.write_text(
        "<!DOCTYPE html><html><head><meta charset='utf-8'><style>"
        "body{font-family:'Times New Roman',Times,serif;font-size:12pt;line-height:1.7;margin:1in;}"
        "h1{font-size:16pt;text-align:center;} h2{font-size:14pt;}</style></head><body>"
        + "\n".join(body)
        + "</body></html>",
        encoding="utf-8",
    )


def print_pdf(html_path: Path, pdf_path: Path):
    profile = Path("/tmp") / f"chrome-pdf-{pdf_path.stem}"
    if profile.exists():
        shutil.rmtree(profile, ignore_errors=True)
    subprocess.run(
        [
            "/opt/google/chrome/chrome",
            "--headless=new",
            "--disable-gpu",
            "--no-sandbox",
            "--disable-dev-shm-usage",
            f"--user-data-dir={profile}",
            "--no-pdf-header-footer",
            "--virtual-time-budget=20000",
            f"--print-to-pdf={pdf_path}",
            html_path.resolve().as_uri(),
        ],
        check=True,
        timeout=90,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def save_tiff(src: Path, dest: Path):
    image = Image.open(src)
    if image.mode != "RGB":
        image = image.convert("RGB")
    dpi = image.info.get("dpi", (400, 400))
    image.save(dest, format="TIFF", compression="tiff_lzw", dpi=dpi)


def export_tables(metrics: dict, folder: Path):
    performance = metrics["performance"]
    with (folder / "Table1_heldout_comparison.csv").open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["System", "ROC-AUC", "ROC-AUC 95% interval", "PR-AUC", "Brier", "Precision at 30%", "Precision at 30% 95% interval", "Recall at 30%", "Alerts at 30%"])
        for row in performance:
            writer.writerow(
                [
                    row["system"],
                    f"{row['roc_auc']:.3f}",
                    f"{row['roc_auc_ci_low']:.3f}–{row['roc_auc_ci_high']:.3f}",
                    f"{row['pr_auc']:.3f}",
                    f"{row['brier']:.3f}",
                    f"{row['precision_at_30']:.3f}",
                    f"{row['precision_at_30_ci_low']:.3f}–{row['precision_at_30_ci_high']:.3f}",
                    f"{row['recall_at_30']:.3f}",
                    row["alerts_at_30"],
                ]
            )
    with (folder / "Table2_paired_tests.csv").open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Comparison", "Metric", "Difference", "95% interval low", "95% interval high", "p-value", "Draws"])
        for row in metrics["paired_tests"]:
            p_value = "<0.001" if row["p_value"] == 0 else f"{row['p_value']:.3f}"
            writer.writerow([row["comparison"], row["metric"], f"{row['difference']:.3f}", f"{row['ci_low']:.3f}", f"{row['ci_high']:.3f}", p_value, row["draws"]])


def supplementary_paragraphs(metrics: dict) -> list[str]:
    paragraphs = [
        "This file accompanies the short communication KB-EWS. It records the checks that are summarised in the main text and the files needed to rebuild them.",
        "# S1. Reproducibility",
        "The program kbews.py, included in this supplement, downloads the public Crunchbase CSV conversion, rebuilds the cohort, revises the rule base, fits the five scores, and writes Table 1, Table 2, Figures 1–3, metrics.json, and warning_cards.csv. The random seed is 7. Raw Crunchbase files are not redistributed. Source: Crunchbase Data Export of 4 December 2015, CSV conversion by Peter Tripp, github.com/notpeter/crunchbase-data, commit be8711fd, licence CC BY-NC 4.0. Commercial use of the extract requires a Crunchbase licence.",
        "# S2. Retention-margin sensitivity",
        "The margin is applied to training-year mean rule strength only. The booster is not refit.",
    ]
    for row in metrics["model"]["retention_margin_sensitivity"]:
        paragraphs.append(
            f"Margin {row['margin']:.2f}: {row['n_rules']} rules retained ({', '.join(row['rules'])}); knowledge-base ROC-AUC {row['roc_auc']:.3f}; precision at 30% {row['precision_at_30']:.3f}."
        )
    paragraphs.append("# S3. Standardized logistic coefficients on the training years")
    paragraphs.append("Positive coefficients raise the log-odds of later closure. Predictors are standardized on 2006–2010. Reference levels are a venture round and the United States.")
    for row in metrics["model"]["logistic_coefficients"]:
        paragraphs.append(f"{row['feature']}: {row['standardized_coef']:.3f}")
    paragraphs.append("# S4. Precision across review budgets")
    paragraphs.append("These are the point estimates drawn in Figure 2. The budget was not chosen on the test labels.")
    for name, series in metrics["budget_profiles"]["systems"].items():
        bits = ", ".join(f"{int(item['rate']*100)}%={item['precision']:.3f}" for item in series)
        paragraphs.append(f"{name}: {bits}")
    paragraphs.append("# S5. Held-out rule firing")
    for row in metrics["rules"]:
        state = "retained" if row["retained"] else "withdrawn"
        paragraphs.append(
            f"{row['code']} {row['name']} ({state}, CF {row['cf']:+.2f}): strength at least 0.5 for {row['fire_rate_failure']:.3f} of later closures and {row['fire_rate_exit']:.3f} of recorded exits."
        )
    return paragraphs


def draw_graphical_abstract(path: Path):
    plt.rcParams.update({"font.family": "DejaVu Serif", "font.size": 11})
    fig = plt.figure(figsize=(13.28, 5.31), dpi=200)
    fig.patch.set_facecolor("white")
    grid = fig.add_gridspec(1, 3, left=0.04, right=0.98, top=0.78, bottom=0.12, wspace=0.28)
    ax0 = fig.add_subplot(grid[0, 0])
    ax1 = fig.add_subplot(grid[0, 1])
    ax2 = fig.add_subplot(grid[0, 2])
    for ax in (ax0, ax1, ax2):
        ax.set_xticks([])
        ax.set_yticks([])
        for spine in ax.spines.values():
            spine.set_visible(False)
    fig.suptitle("Early warning at the first priced round of a digital venture", fontsize=15, x=0.04, ha="left", color="#12344D")
    ax0.set_xlim(0, 1)
    ax0.set_ylim(0, 1)
    ax0.text(0, 0.92, "Decision time", fontsize=13, color="#12344D", fontweight="regular")
    ax0.text(0, 0.72, "First priced round only", fontsize=12, color="#1c2430")
    ax0.text(0, 0.48, "Allowed\nlog capital, days since founding,\nscope, digital depth, ecosystem,\nround type, country", fontsize=10, color="#1c2430", linespacing=1.45, va="top")
    ax0.text(0, 0.16, "Refused in code\ntotal capital, later rounds,\nlast-round date, status", fontsize=10, color="#7A2E2E", linespacing=1.45, va="center")

    ax1.barh([0, 1], [0.510, 0.760], color=["#94A3B8", "#1F4E79"], height=0.55)
    ax1.set_yticks([0, 1])
    ax1.set_yticklabels(["Base rate", "Top 30% precision"])
    ax1.set_xlim(0, 1)
    ax1.set_xlabel("Closure share among reviewed ventures")
    ax1.set_title("Held-out screen, n = 737", loc="left", fontsize=13, color="#12344D", pad=8)
    ax1.spines["left"].set_visible(True)
    ax1.spines["bottom"].set_visible(True)
    ax1.tick_params(axis="y", length=0, labelsize=10)
    ax1.text(0.760, 1, "  0.760", va="center", fontsize=10, color="#1F4E79")
    ax1.text(0.510, 0, "  0.510", va="center", fontsize=10, color="#475569")

    ax2.set_xlim(0, 1)
    ax2.set_ylim(0, 1)
    ax2.text(0, 0.92, "What stays on the card", fontsize=13, color="#12344D")
    ax2.text(
        0,
        0.62,
        "R1 Resource strain\nR2 Unfunded scope\nR3 Shallow digital base\nR6 Resourced focus",
        fontsize=11,
        color="#1c2430",
        linespacing=1.55,
        va="top",
    )
    ax2.text(0, 0.12, "The rules name the mechanism.\nThey do not replace the rank.", fontsize=11, color="#12344D", va="center")
    fig.savefig(path, dpi=200, facecolor="white")
    plt.close(fig)


def main():
    for item in HIGHLIGHTS:
        if len(item) > 85:
            raise SystemExit(f"Highlight exceeds 85 characters: {item}")
    folders = ensure_dirs()
    blocks = parse_markdown(MANUSCRIPT.read_text(encoding="utf-8"))
    metrics = json.loads((ROOT / "outputs" / "metrics.json").read_text(encoding="utf-8"))

    docx_path = folders["01_manuscript"] / "KB-EWS_manuscript.docx"
    html_path = folders["01_manuscript"] / "KB-EWS_manuscript.html"
    pdf_path = folders["01_manuscript"] / "KB-EWS_manuscript.pdf"
    build_manuscript_docx(blocks, docx_path)
    build_manuscript_html(blocks, html_path)
    print_pdf(html_path, pdf_path)

    figure_map = {
        "fig1_roc.png": ("Figure1_ROC.png", "Figure1_ROC.tiff"),
        "fig2_budget.png": ("Figure2_precision_budget.png", "Figure2_precision_budget.tiff"),
        "fig3_rules.png": ("Figure3_rule_firing.png", "Figure3_rule_firing.tiff"),
    }
    for source_name, (png_name, tiff_name) in figure_map.items():
        source = ROOT / "outputs" / source_name
        shutil.copy2(source, folders["02_figures"] / png_name)
        save_tiff(source, folders["02_figures"] / tiff_name)

    export_tables(metrics, folders["03_tables"])
    shutil.copy2(ROOT / "outputs" / "metrics.json", folders["04_supporting_data"] / "metrics.json")
    shutil.copy2(ROOT / "outputs" / "warning_cards.csv", folders["04_supporting_data"] / "warning_cards.csv")
    (folders["04_supporting_data"] / "DATA_STATEMENT.txt").write_text(
        "\n".join(
            [
                "KB-EWS supporting data",
                "",
                "metrics.json is the numeric source for Table 1, Table 2, the sensitivity check, the logistic coefficients, and the rule-firing rates.",
                "warning_cards.csv has one row for each of the 737 ventures first funded in 2012–2013.",
                "The reviewer fields are closure_probability, queue_rank, in_review_queue, and fired_rules.",
                "later_closed is the held-out outcome used to audit the queue. It is not a model input.",
                "",
                "The raw source is the Crunchbase Data Export of 4 December 2015, redistributed as CSV by Peter Tripp",
                "(github.com/notpeter/crunchbase-data, commit be8711fd) under CC BY-NC 4.0.",
                "Attribute Crunchbase. The raw company and round files are not included in this package.",
                "Commercial use of that extract requires a separate Crunchbase licence.",
                "Running python kbews.py with seed 7 rebuilds these files from that public export.",
                "",
            ]
        ),
        encoding="utf-8",
    )

    abstract_png = folders["05_graphical_abstract"] / "graphical_abstract.png"
    draw_graphical_abstract(abstract_png)
    save_tiff(abstract_png, folders["05_graphical_abstract"] / "graphical_abstract.tiff")

    supplement = supplementary_paragraphs(metrics)
    supp_docx = folders["06_supplementary"] / "Supplementary_material.docx"
    supp_html = folders["06_supplementary"] / "Supplementary_material.html"
    write_simple_docx(supp_docx, "KB-EWS supplementary material", supplement)
    write_simple_html(supp_html, "KB-EWS supplementary material", supplement)
    print_pdf(supp_html, folders["06_supplementary"] / "Supplementary_material.pdf")
    shutil.copy2(ROOT / "kbews.py", folders["06_supplementary"] / "kbews.py")

    cover = [
        "Dear Editor-in-Chief,",
        "Please consider the enclosed short communication, “A Knowledge-Based Early-Warning System for Entrepreneurial Failure during Digital Transformation,” for Knowledge-Based Systems.",
        "The manuscript presents KB-EWS, a warning tool for the first priced round of a digital venture. The program admits only signals known at that round and stops if lifetime funding fields enter the design matrix. Seven certainty-factor rules were written before estimation. Four survive a retention test on ventures first funded in 2006–2010. On 737 ventures first funded in 2012–2013, the pre-specified depth-3 booster has ROC area 0.740. Reviewing the highest-scoring 30% gives precision 0.760 against a base rate of 0.510. Paired tests do not separate the booster from logistic regression. Stacking the certainty factor into the probability does not improve the review queue. The retained rules remain on the warning card.",
        "The work is original, is not under review elsewhere, and has not been published. The author has no competing interest to declare. AI-assisted tools were used for language editing; the author reviewed the manuscript and takes responsibility for it. The tables and figures were generated from the accompanying code.",
        "Corresponding author: Balavaraprasad Reddy Yenugula, Doctor of Business Administration candidate.",
        "Yours sincerely,",
        "Balavaraprasad Reddy Yenugula",
    ]
    cover_docx = folders["07_cover_letter_and_highlights"] / "Cover_letter.docx"
    cover_html = folders["07_cover_letter_and_highlights"] / "Cover_letter.html"
    write_simple_docx(cover_docx, "Cover letter", cover)
    write_simple_html(cover_html, "Cover letter", cover)
    print_pdf(cover_html, folders["07_cover_letter_and_highlights"] / "Cover_letter.pdf")
    (folders["07_cover_letter_and_highlights"] / "Highlights.txt").write_text(
        "Paste these five highlights into the Elsevier highlights boxes. Each is within the 85-character limit.\n\n"
        + "\n".join(f"{i}. {item}  ({len(item)} characters)" for i, item in enumerate(HIGHLIGHTS, 1))
        + "\n",
        encoding="utf-8",
    )

    (SUB / "UPLOAD_LIST.txt").write_text(
        "\n".join(
            [
                "Knowledge-Based Systems submission files",
                "Portal step: submission-files-upload",
                "",
                "Manuscript: 01_manuscript/KB-EWS_manuscript.docx",
                "PDF copy of the same manuscript: 01_manuscript/KB-EWS_manuscript.pdf",
                "Figures: 02_figures/Figure1_ROC.tiff, Figure2_precision_budget.tiff, Figure3_rule_firing.tiff",
                "Tables: 03_tables/Table1_heldout_comparison.csv, Table2_paired_tests.csv",
                "Supporting data: 04_supporting_data/metrics.json, warning_cards.csv, DATA_STATEMENT.txt",
                "Graphical abstract: 05_graphical_abstract/graphical_abstract.tiff",
                "Supplementary material: 06_supplementary/Supplementary_material.pdf and kbews.py",
                "Cover letter: 07_cover_letter_and_highlights/Cover_letter.docx",
                "Highlights to paste into the form: 07_cover_letter_and_highlights/Highlights.txt",
                "",
                "The raw Crunchbase CSV files are not included. See DATA_STATEMENT.txt.",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print("package written to", SUB)


if __name__ == "__main__":
    main()
